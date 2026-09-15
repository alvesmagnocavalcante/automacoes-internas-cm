"""Leitura e validação dos relatórios OPERA, CMFlex e Rede."""

from __future__ import annotations

import re
import warnings
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook

from automations.recebimentos.models import CmflexPayment, OperaPayment, RedePayment
from automations.recebimentos.normalization import normalize


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def opera_xml_has_transactions(path: Path) -> bool:
    if path.suffix.casefold() != ".xml":
        return True
    try:
        root = ElementTree.parse(path).getroot()
    except ElementTree.ParseError:
        return False
    return any(_local_name(node.tag) == "G_TRANSACTION" for node in root.iter())


def identifier(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def decimal_value(value: object, *, field: str, row_number: int) -> Decimal:
    if isinstance(value, Decimal):
        return value
    text = (
        ("" if value is None else str(value)).strip().replace("R$", "").replace(" ", "")
    )
    if not text:
        raise ValueError(f"Campo {field!r} vazio na linha {row_number}.")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as error:
        raise ValueError(
            f"Valor inválido em {field!r} na linha {row_number}: {value!r}."
        ) from error


def optional_decimal_value(value: object, *, field: str, row_number: int) -> Decimal:
    if value is None or not str(value).strip():
        return Decimal("0")
    return decimal_value(value, field=field, row_number=row_number)


def card_last_four(*values: object) -> str:
    for value in values:
        digits = re.sub(r"\D", "", str(value or ""))
        if len(digits) >= 4:
            return digits[-4:]
    return ""


def _worksheet_rows(
    path: Path, required_headers: set[str]
) -> tuple[list[str], list[tuple[object, ...]], int]:
    if not path.is_file():
        raise FileNotFoundError(f"Relatório não encontrado: {path}")
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        workbook = load_workbook(path, data_only=True, read_only=False)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    for index, row in enumerate(rows[:20], start=1):
        headers = [normalize(value) for value in row]
        if required_headers.issubset(headers):
            return headers, rows[index:], index
    raise ValueError(
        f"Cabeçalho obrigatório não encontrado em {path.name}: "
        + ", ".join(sorted(required_headers))
    )


def _records(
    path: Path, required_headers: set[str]
) -> list[tuple[int, dict[str, object]]]:
    headers, rows, header_row = _worksheet_rows(path, required_headers)
    return [
        (row_number, dict(zip(headers, row, strict=False)))
        for row_number, row in enumerate(rows, start=header_row + 1)
        if any(value is not None for value in row)
    ]


def parse_cmflex(path: Path) -> list[CmflexPayment]:
    required = {
        "cliente",
        "numero",
        "portadorforma",
        "tipodedocumento",
        "valor",
    }
    payments = []
    for row_number, row in _records(path, required):
        payments.append(
            CmflexPayment(
                row_number=row_number,
                document_number=identifier(row["numero"]),
                customer=identifier(row["cliente"]),
                carrier=identifier(row["portadorforma"]),
                document_type=identifier(row["tipodedocumento"]),
                amount=optional_decimal_value(
                    row["valor"], field="Valor", row_number=row_number
                ),
            )
        )
    return payments


def parse_rede(path: Path) -> list[RedePayment]:
    required = {
        "status da venda",
        "valor da venda original",
        "modalidade",
        "bandeira",
        "nsu/cv",
        "numero da autorizacao (auto)",
        "numero do cartao",
        "id carteira digital",
    }
    payments = []
    for row_number, row in _records(path, required):
        payments.append(
            RedePayment(
                row_number=row_number,
                status=identifier(row["status da venda"]),
                modality=identifier(row["modalidade"]),
                brand=identifier(row["bandeira"]),
                card_last_four=card_last_four(
                    row["numero do cartao"], row["id carteira digital"]
                ),
                nsu=identifier(row["nsu/cv"]),
                amount=decimal_value(
                    row.get("valor da venda atualizado")
                    if row.get("valor da venda atualizado") not in (None, "", "-")
                    else row["valor da venda original"],
                    field="valor da venda",
                    row_number=row_number,
                ),
            )
        )
    return payments


def _opera_payment(row: dict[str, object], row_number: int) -> OperaPayment:
    return OperaPayment(
        transaction_id=identifier(row["trx_no"]),
        folio_number=identifier(row.get("folio_no")),
        transaction_code=identifier(row.get("trx_code") or row.get("grp_first")),
        description=identifier(row["trx_desc"]),
        card_last_four=card_last_four(row.get("card_number")),
        amount=-decimal_value(
            row["guest_account_credit"],
            field="GUEST_ACCOUNT_CREDIT",
            row_number=row_number,
        ),
    )


def _parse_opera_xml(path: Path) -> list[OperaPayment]:
    if not path.is_file():
        raise FileNotFoundError(f"Relatório não encontrado: {path}")
    try:
        root = ElementTree.parse(path).getroot()
    except ElementTree.ParseError as error:
        raise ValueError(f"XML inválido em {path}: {error}") from error
    payments = []
    required = {"trx_no", "trx_code", "trx_desc", "guest_account_credit"}
    transaction_nodes = (
        node for node in root.iter() if _local_name(node.tag) == "G_TRANSACTION"
    )
    for row_number, node in enumerate(transaction_nodes, start=1):
        row = {
            normalize(_local_name(child.tag)): (child.text or "").strip()
            for child in node
        }
        missing = required - row.keys()
        if missing:
            raise ValueError(
                f"Campos ausentes na transação OPERA {row_number}: "
                + ", ".join(sorted(missing))
            )
        payments.append(_opera_payment(row, row_number))
    if not payments:
        raise ValueError(f"Nenhuma transação encontrada no XML do OPERA: {path}")
    return payments


def _parse_opera_xlsx(path: Path) -> list[OperaPayment]:
    required = {"trx_no", "trx_desc", "guest_account_credit"}
    payments = [
        _opera_payment(row, row_number) for row_number, row in _records(path, required)
    ]
    if not payments:
        raise ValueError(f"Nenhuma transação encontrada no XLSX do OPERA: {path}")
    return payments


def parse_opera(path: Path) -> list[OperaPayment]:
    suffix = path.suffix.casefold()
    if suffix == ".xml":
        return _parse_opera_xml(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_opera_xlsx(path)
    raise ValueError(
        f"Formato do relatório OPERA não suportado: {path.suffix or 'sem extensão'}. "
        "Use XML ou XLSX."
    )
