from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from openpyxl import Workbook

from automations.recebimentos.parsers import (
    card_last_four,
    decimal_value,
    optional_decimal_value,
    parse_opera,
    parse_rede,
)


class ParserTests(TestCase):
    def test_extracts_card_last_four_from_rede_columns(self):
        self.assertEqual(card_last_four("-", "443435******3935"), "3935")

    def test_parses_brazilian_currency_exactly(self):
        self.assertEqual(
            decimal_value("R$ 1.234,56", field="valor", row_number=2),
            Decimal("1234.56"),
        )

    def test_empty_optional_currency_is_zero(self):
        self.assertEqual(
            optional_decimal_value(None, field="valor", row_number=2), Decimal("0")
        )

    def test_numeric_zero_is_not_treated_as_empty(self):
        self.assertEqual(decimal_value(0, field="valor", row_number=2), Decimal("0"))

    def test_opera_payment_sign_is_normalized_and_refund_becomes_negative(self):
        xml = """<?xml version="1.0" encoding="UTF-8"?>
<FINPAYMENTS><G_TRANSACTION>
<TRX_NO>123</TRX_NO><FOLIO_NO>456</FOLIO_NO><TRX_CODE>9302</TRX_CODE>
<TRX_DESC>Rede Visa Credito</TRX_DESC><CARD_NUMBER>3935</CARD_NUMBER>
<GUEST_ACCOUNT_CREDIT>5.50</GUEST_ACCOUNT_CREDIT>
</G_TRANSACTION></FINPAYMENTS>"""
        with TemporaryDirectory() as directory:
            path = Path(directory) / "opera.xml"
            path.write_text(xml, encoding="utf-8")

            payments = parse_opera(path)

        self.assertEqual(payments[0].amount, Decimal("-5.50"))
        self.assertEqual(payments[0].card_last_four, "3935")

    def test_parses_opera_xlsx_used_by_historical_report(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "opera.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(
                [
                    "TRX_NO",
                    "FOLIO_NO",
                    "TRX_CODE",
                    "TRX_DESC",
                    "CARD_NUMBER",
                    "GUEST_ACCOUNT_CREDIT",
                ]
            )
            sheet.append([123, 456, 9302, "Rede Visa Crédito", "**** 3223", -81])
            workbook.save(path)

            payments = parse_opera(path)

        self.assertEqual(len(payments), 1)
        self.assertEqual(payments[0].transaction_id, "123")
        self.assertEqual(payments[0].folio_number, "456")
        self.assertEqual(payments[0].card_last_four, "3223")
        self.assertEqual(payments[0].amount, Decimal("81"))

    def test_opera_xlsx_accepts_grp_first_as_transaction_code(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "opera.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(
                ["GRP_FIRST", "TRX_NO", "TRX_DESC", "GUEST_ACCOUNT_CREDIT"]
            )
            sheet.append([9086, 123, "Dinheiro", -10])
            workbook.save(path)

            payments = parse_opera(path)

        self.assertEqual(payments[0].transaction_code, "9086")

    def test_rede_accepts_historical_file_without_updated_amount(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "rede.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["EXTRATO"])
            sheet.append(
                [
                    "status da venda",
                    "valor da venda original",
                    "modalidade",
                    "bandeira",
                    "NSU/CV",
                    "número da autorização (Auto)",
                    "número do cartão",
                    "id carteira digital",
                ]
            )
            sheet.append(
                ["aprovada", 20, "crédito", "Visa", "1", "2", "****1234", "-"]
            )
            workbook.save(path)

            payments = parse_rede(path)

        self.assertEqual(payments[0].amount, Decimal("20"))
