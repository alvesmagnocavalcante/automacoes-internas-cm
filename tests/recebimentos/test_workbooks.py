import os
import shutil
import subprocess
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest import TestCase, skipUnless
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from automations.recebimentos.domain import reconcile
from automations.recebimentos.parsers import parse_cmflex, parse_opera, parse_rede
from automations.recebimentos.workbooks import (
    _copy_with_permissions,
    _reset_inherited_permissions,
    save_conference_workbooks,
)


class WorkbookPermissionsTests(TestCase):
    def test_windows_resets_acl_after_copy(self):
        with TemporaryDirectory() as directory:
            src = Path(directory) / "origem.xlsx"
            dst = Path(directory) / "relatório final.xlsx"
            src.write_bytes(b"workbook")

            def check_copy(*args, **kwargs):
                self.assertEqual(dst.read_bytes(), b"workbook")

            with (
                patch("automations.recebimentos.workbooks.os", SimpleNamespace(name="nt")),
                patch("automations.recebimentos.workbooks.subprocess.run", side_effect=check_copy) as run,
            ):
                _copy_with_permissions(src, dst)
                run.assert_called_once_with(
                    ["icacls", str(dst), "/reset", "/Q"],
                    check=True, capture_output=True,
                )

    def test_windows_acl_failure_is_propagated(self):
        path = Path("report.xlsx")
        with (
            patch("automations.recebimentos.workbooks.os", SimpleNamespace(name="nt")),
            patch("automations.recebimentos.workbooks.subprocess.run") as run,
        ):
            run.side_effect = subprocess.CalledProcessError(1, "icacls")
            with self.assertRaises(subprocess.CalledProcessError):
                _reset_inherited_permissions(path)

    def test_posix_applies_file_mode_after_copy(self):
        with TemporaryDirectory() as directory:
            src = Path(directory) / "source.xlsx"
            dst = Path(directory) / "target.xlsx"
            src.write_bytes(b"workbook")
            with patch("automations.recebimentos.workbooks.os") as operating_system:
                operating_system.name = "posix"
                operating_system.chmod.side_effect = lambda *args: self.assertEqual(
                    dst.read_bytes(), b"workbook"
                )
                _copy_with_permissions(src, dst)
                operating_system.chmod.assert_called_once_with(dst, 0o644)

    @skipUnless(os.name == "posix", "Requer permissões POSIX reais")
    def test_posix_overrides_restrictive_umask(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            src = root / "source.xlsx"
            src.write_bytes(b"workbook")
            previous_umask = os.umask(0o077)
            try:
                destination = root / "reports"
                destination.mkdir()
                dst = destination / "target.xlsx"
                _copy_with_permissions(src, dst)
            finally:
                os.umask(previous_umask)
            self.assertEqual(dst.stat().st_mode & 0o777, 0o644)


class ConferenceWorkbooksTests(TestCase):
    def test_separates_companies_without_overwriting_previous_result(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            opera, cmflex, rede = (root / name for name in ("opera.xml", "cmflex.xlsx", "rede.xlsx"))
            self._write_opera(opera)
            self._write_cmflex(cmflex)
            self._write_rede(rede)
            result = reconcile(parse_opera(opera), parse_cmflex(cmflex), parse_rede(rede))
            archive = root / "conferencias"
            charme, _ = save_conference_workbooks(
                opera, cmflex, rede, result, archive, date(2026, 8, 1),
                "CHARME", "CHARME",
            )
            magna, _ = save_conference_workbooks(
                opera, cmflex, rede, result, archive, date(2026, 8, 1),
                "MAGNA", "MAGNA",
            )
            self.assertEqual(charme.parent, magna.parent)
            self.assertEqual(len(list(charme.glob("*.xlsx"))), 3)
            self.assertEqual(len(list(magna.glob("*.xlsx"))), 3)

    def test_creates_only_three_marked_excel_files(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            opera = root / "opera.xml"
            cmflex = root / "cmflex.xlsx"
            rede = root / "rede.xlsx"
            self._write_opera(opera)
            self._write_cmflex(cmflex)
            self._write_rede(rede)
            result = reconcile(
                parse_opera(opera), parse_cmflex(cmflex), parse_rede(rede)
            )
            destination = root / "conferencias/08 - AGOSTO/01"
            destination.mkdir(parents=True)
            (destination / "Conferencia antiga.json").write_text("{}")

            with patch(
                "automations.recebimentos.workbooks.shutil.copyfile",
                wraps=shutil.copyfile,
            ) as copyfile:
                saved_directory, saved = save_conference_workbooks(
                    opera,
                    cmflex,
                    rede,
                    result,
                    root / "conferencias",
                    date(2026, 8, 1),
                    "MAGNA - Magna Praia Hotel",
                )

            self.assertEqual(saved_directory, destination)
            self.assertEqual(
                sorted(path.name for path in destination.iterdir()),
                [
                    "CmFlex Magna 01.08.xlsx",
                    "Opera Magna 01.08.xlsx",
                    "Rede Magna 01.08.xlsx",
                ],
            )
            self.assertEqual(set(saved), {"Opera", "CmFlex", "Rede"})
            self.assertEqual(copyfile.call_count, 3)
            self.assertEqual(
                {call.args[1] for call in copyfile.call_args_list},
                set(saved.values()),
            )

            opera_book = load_workbook(saved["Opera"])
            self.assertEqual(opera_book.active["F2"].fill.fgColor.rgb, "00FFFF00")
            self.assertEqual(opera_book.active["F4"].fill.fgColor.rgb, "00FF99CC")
            self.assertEqual(opera_book.active["F5"].fill.fgColor.rgb, "00FF99CC")
            opera_book.close()

            cmflex_book = load_workbook(saved["CmFlex"])
            self.assertEqual(cmflex_book.active.max_column, 9)
            self.assertEqual(cmflex_book.active["D2"].fill.fgColor.rgb, "00FFFF00")
            cmflex_book.close()

            rede_book = load_workbook(saved["Rede"])
            self.assertEqual(rede_book.active["D3"].fill.fgColor.rgb, "00FF99CC")
            rede_book.close()

    @staticmethod
    def _write_opera(path: Path) -> None:
        path.write_text(
            """<FINPAYMENTS>
<G_TRANSACTION><TRX_NO>cash-1</TRX_NO><TRX_CODE>9086</TRX_CODE>
<TRX_DESC>Dinheiro</TRX_DESC><GUEST_ACCOUNT_CREDIT>-2</GUEST_ACCOUNT_CREDIT>
</G_TRANSACTION>
<G_TRANSACTION><TRX_NO>pix-1</TRX_NO><TRX_CODE>9087</TRX_CODE>
<TRX_DESC>Pix</TRX_DESC><GUEST_ACCOUNT_CREDIT>-6</GUEST_ACCOUNT_CREDIT>
</G_TRANSACTION>
<G_TRANSACTION><TRX_NO>card-1</TRX_NO><TRX_CODE>9302</TRX_CODE>
<TRX_DESC>Rede Visa Credito</TRX_DESC><CARD_NUMBER>1234</CARD_NUMBER>
<GUEST_ACCOUNT_CREDIT>-4</GUEST_ACCOUNT_CREDIT></G_TRANSACTION>
<G_TRANSACTION><TRX_NO>card-2</TRX_NO><TRX_CODE>9302</TRX_CODE>
<TRX_DESC>Rede Visa Credito</TRX_DESC><CARD_NUMBER>1234</CARD_NUMBER>
<GUEST_ACCOUNT_CREDIT>-12</GUEST_ACCOUNT_CREDIT></G_TRANSACTION>
</FINPAYMENTS>""",
            encoding="utf-8",
        )

    @staticmethod
    def _write_cmflex(path: Path) -> None:
        book = Workbook()
        sheet = book.active
        sheet.title = "LancamentoDeDocumentoCAR"
        sheet.append(
            [
                "Cliente",
                "Numero",
                "PortadorForma",
                "TipoDeDocumento",
                "SistemaLancamento",
                "Valor",
                "ValorOutraMoeda",
                "Saldo",
                "ValorAlteradores",
                "ValorLiquido",
                "DataEmissao",
            ]
        )
        sheet.append(
            [
                "DINHEIRO",
                "cash-1",
                "CAIXA",
                "MOVIMENTO DE CAIXA",
                "Integração Back-Office",
                2,
                0,
                2,
                0,
                2,
                "01/08/2026",
            ]
        )
        sheet.append(
            [
                "DEPOSITO",
                "pix-1",
                "CONTA",
                "Deposito",
                "Integração Back-Office",
                6,
                0,
                6,
                0,
                6,
                "01/08/2026",
            ]
        )
        book.save(path)
        book.close()

    @staticmethod
    def _write_rede(path: Path) -> None:
        book = Workbook()
        sheet = book.active
        sheet.title = "vendas"
        sheet.append(["EXTRATO PARA SIMPLES CONFERÊNCIA"])
        sheet.append(
            [
                "data da venda",
                "hora da venda",
                "status da venda",
                "valor da venda original",
                "modalidade",
                "bandeira",
                "NSU/CV",
                "número da autorização (Auto)",
                "número do cartão",
                "id carteira digital",
            ]
        )
        sheet.append(
            [
                "01/08/2026",
                "10:00:00",
                "aprovada",
                16,
                "crédito",
                "Visa",
                "nsu-1",
                "auth-1",
                "****1234",
                "-",
            ]
        )
        book.save(path)
        book.close()
